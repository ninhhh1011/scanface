"""Synthetic algorithm/protocol tests; never evidence of human face accuracy."""
import base64
import io
import time

import numpy as np
import pytest
from PIL import Image
from cryptography.exceptions import InvalidTag

from app.security import sign_scope, verify_scope, encrypt_template, decrypt_template
from app.face import match_identity, check_liveness, decode_frame, normalize, FaceError
from app.ingestion import extract, chunks, validate_embeddings, IngestionError


def test_scope_expiry_signature_and_body_binding():
    now = int(time.time())
    scope = dict(actor_id="a", session_id="s", employee_id="e", action="ENROLL", exp=now+30, nonce="unique-nonce")
    token = sign_scope(scope, "k" * 32)
    assert verify_scope(token, "k" * 32, {"employee_id": "e"}) == scope
    for key, expected in [("employee_id", "other"), ("action", "CHECK_IN")]:
        with pytest.raises(ValueError): verify_scope(token, "k" * 32, {key: expected})
    with pytest.raises(ValueError): verify_scope(token, "z" * 32, {})
    scope["exp"] = now-1
    with pytest.raises(ValueError): verify_scope(sign_scope(scope, "k" * 32), "k" * 32, {})


def test_authenticated_template_context_and_roundtrip():
    key = bytes(range(32))
    vector = normalize(np.arange(1, 129, dtype=np.float32))
    blob = encrypt_template(vector, key, "employee-a", "sface-v1")
    assert np.allclose(decrypt_template(blob, key, "employee-a", "sface-v1", 128), vector)
    assert blob != encrypt_template(vector, key, "employee-a", "sface-v1")
    with pytest.raises(InvalidTag): decrypt_template(blob, key, "employee-b", "sface-v1", 128)
    with pytest.raises(InvalidTag): decrypt_template(blob[:-1] + bytes([blob[-1] ^ 1]), key, "employee-a", "sface-v1", 128)


def test_match_groups_templates_per_identity_and_rejects_unknown_ambiguity():
    q = np.array([1., 0.])
    assert match_identity(q, [], .5, .1) == ("NO_ENROLLED_PROFILES", None)
    assert match_identity(q, [("a", np.array([0., 1.]))], .5, .1) == ("UNKNOWN_PERSON", None)
    assert match_identity(q, [("a", q), ("a", q)], .5, .1) == ("MATCHED", "a")
    assert match_identity(q, [("a", q), ("b", normalize(np.array([1., .1])))], .5, .1) == ("AMBIGUOUS_MATCH", None)


def test_liveness_requires_order_motion_return_and_timing():
    stamps = [1000, 1600, 2200, 2800, 3400]
    assert check_liveness([0., .10, .25, .10, 0.], stamps, "LEFT")
    assert not check_liveness([0.] * 5, stamps, "LEFT")
    assert not check_liveness([0., .10, .25, .10, 0.], stamps, "RIGHT")
    assert not check_liveness([0., -.10, -.25, -.10, 0.], [1000]*5, "LEFT")
    assert not check_liveness([0., -.10, -.25, -.25, -.25], stamps, "LEFT")


def test_jpeg_bounds_and_bad_payload():
    buf = io.BytesIO()
    Image.new("RGB", (320, 240)).save(buf, "JPEG")
    assert decode_frame(base64.b64encode(buf.getvalue()).decode()).shape == (240, 320, 3)
    for value in ["not base64", base64.b64encode(b"not an image").decode()]:
        with pytest.raises(FaceError): decode_frame(value)
    buf = io.BytesIO()
    Image.new("RGB", (2100, 200)).save(buf, "JPEG")
    with pytest.raises(FaceError): decode_frame(base64.b64encode(buf.getvalue()).decode())


def test_extract_markdown_chunks_preserve_metadata_and_no_empty_success():
    sections = extract("# Nghỉ phép\n\nNhân viên gửi đơn.\n".encode(), "policy.md", "text/markdown")
    result = chunks(sections)
    assert result[0]["section"] == "Nghỉ phép"
    assert "Nhân viên" in result[0]["content"]
    assert result[0]["page"] is None
    with pytest.raises(IngestionError): extract(b"  ", "empty.txt", "text/plain")
    long = chunks([{"content": "x" * 4000, "section": "a", "page": 1}])
    assert len(long) > 1 and all(len(c["content"]) <= 1200 for c in long)


def test_embedding_response_dimension_nan_and_missing_rows():
    assert validate_embeddings({"data": [{"index": 0, "embedding": [1., 2.]}]}, 1, 2) == [[1., 2.]]
    for payload in [{"data": []}, {"data": [{"index": 0, "embedding": [1.]}]}, {"data": [{"index": 0, "embedding": [float("nan"), 1.]}]}]:
        with pytest.raises(IngestionError): validate_embeddings(payload, 1, 2)


def test_docx_heading_table_order_and_scan_pdf_rejection():
    from docx import Document
    from pypdf import PdfWriter
    document = Document()
    document.add_heading("Section one", 1)
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table belongs to section one"
    document.add_heading("Section two", 1)
    document.add_paragraph("After table")
    buffer = io.BytesIO()
    document.save(buffer)
    parts = extract(buffer.getvalue(), "policy.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "Table belongs" in parts[0]["content"] and parts[0]["section"] == "Section one"
    assert "Table belongs" not in parts[1]["content"]
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    pdf.write(buffer)
    with pytest.raises(IngestionError, match="UNSUPPORTED_SCAN_NEEDS_TEXT"):
        extract(buffer.getvalue(), "scan.pdf", "application/pdf")
